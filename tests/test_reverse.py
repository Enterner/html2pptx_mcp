"""逆转管线单测：outline / 选择性逆转 / ea 字体槽优先 / 占位模式。

输入 pptx 一律在测试内由样例 HTML 编译或 python-pptx 构造，不依赖任何
仓库外的二进制夹具，CI（无 PowerPoint）可跑。
"""

from __future__ import annotations

from pathlib import Path

from html2pptx_mcp.convert import CompileOptions, compile_html
from html2pptx_mcp.convert.reverse import outline_pptx, reverse_pptx

EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "html" / "01_title.html"


def _compiled_pptx(tmp_path: Path) -> Path:
    result = compile_html(
        EXAMPLE.read_text(encoding="utf-8"), options=CompileOptions(base_dir=EXAMPLE.parent)
    )
    assert result.ok
    pptx_path = tmp_path / "example.pptx"
    result.save(pptx_path)
    return pptx_path


def test_outline_roles_and_theme(tmp_path: Path) -> None:
    pptx = _compiled_pptx(tmp_path)
    outline = outline_pptx(pptx)
    assert outline["pages"] == 1
    assert outline["canvas"] == {"w": 1280, "h": 720}
    page = outline["inventory"][0]
    assert page["page"] == 1
    roles = page["roles"]
    assert roles.get("title") == 1
    assert set(outline["theme"]["fonts"].values()) - {None}


def test_reverse_full_roundtrip_compiles(tmp_path: Path) -> None:
    """逆转 → 重编译的闭环必须成立：样例逆转出的 HTML 是合法契约子集。"""
    pptx = _compiled_pptx(tmp_path)
    result = reverse_pptx(pptx, tmp_path / "reversed")
    html_file = result.html_paths[0]
    recompiled = compile_html(
        html_file.read_text(encoding="utf-8"), options=CompileOptions(base_dir=html_file.parent)
    )
    assert recompiled.ok
    assert recompiled.report.stats["slides"] == 1


def test_reverse_element_selection(tmp_path: Path) -> None:
    """elements 只保留选中的顶层形状；未选中者（含图片）不进片段。

    片段含根 div（.stage 内容）本身，故 div 数 = 选中数 + 1。
    """
    pptx = _compiled_pptx(tmp_path)
    result = reverse_pptx(pptx, tmp_path / "sel", pages=[1], elements={1: [0]})
    fragment = result.html_by_page[1]
    assert fragment.count("<div") == 2  # 根 + 选中的 0 号形状
    assert "assets/" not in fragment  # 样例页 0 号形状不是图片


def test_reverse_placeholder_text(tmp_path: Path) -> None:
    pptx = _compiled_pptx(tmp_path)
    result = reverse_pptx(pptx, tmp_path / "ph", placeholder_text=True)
    fragment = result.html_by_page[1]
    original_title = "标题"  # 样例标题含「标题」二字以外的具体文案
    assert fragment  # 片段非空
    # 占位模式：原文案不出现，等长全角占位出现
    assert "示例" not in fragment and "文" in fragment
    del original_title


def test_reverse_ea_font_priority(tmp_path: Path) -> None:
    """回归：CJK 字形由 ea 槽决定——run 只写 a:ea 时逆转必须取 ea 的字体名。

    原始 bug：只读 a:latin，副标题的 ea=华康宋体W12(P) 被漏掉，重建走
    等线（真字体真粗体），与原版的回退渲染字重不一致。
    """
    from pptx import Presentation
    from pptx.util import Emu

    pptx_path = tmp_path / "ea.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(914400), Emu(914400))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "副标题文案"
    # python-pptx 的 font.name 只写 a:latin；ea 需要手工补
    run.font.name = "Arial"
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.makeelement(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}ea",
        {"typeface": "华康宋体W12(P)"},
    )
    r_pr.append(ea)
    presentation.save(pptx_path)

    result = reverse_pptx(pptx_path, tmp_path / "rev")
    fragment = result.html_by_page[1]
    assert "华康宋体W12(P)" in fragment


def test_reverse_notes_are_namespaced(tmp_path: Path) -> None:
    """逆转的降级记录走自有 Note 通道（「域.规则」命名），不混入契约诊断码。"""
    pptx = _compiled_pptx(tmp_path)
    result = reverse_pptx(pptx, tmp_path / "notes")
    for note in result.notes:
        assert "." in note.rule, f"降级规则 {note.rule!r} 应带命名空间"
        assert not note.rule.isupper(), (
            f"降级规则 {note.rule!r} 长得像诊断码，应走 Note 通道而非契约诊断"
        )

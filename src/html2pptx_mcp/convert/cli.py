"""命令行入口：html2pptx-cli validate|build|contract|inspect。

退出码固定三档，便于脚本与 CI 判定：

- ``0`` 通过；
- ``2`` 存在契约错误（``--strict`` 时警告也算）；
- ``1`` 内部错误（文件读不了、pptx 打不开等）。
"""

from __future__ import annotations

import argparse
import json
import sys
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from .api import CompileOptions, compile_html, validate_html
from .contract import load_contract
from .contract_docs import render_json, render_llm, render_markdown
from .diagnostics import Report

EXIT_OK = 0
EXIT_INTERNAL = 1
EXIT_CONTRACT = 2


def main(argv: Sequence[str] | None = None) -> int:
    parser = _build_parser()
    args = parser.parse_args(argv)
    try:
        if args.command == "validate":
            return _validate(args)
        if args.command == "build":
            return _build(args)
        if args.command == "contract":
            return _contract(args)
        return _inspect(args)
    except OSError as error:
        print(f"文件错误：{error}", file=sys.stderr)
        return EXIT_INTERNAL
    except ValueError as error:
        print(f"参数错误：{error}", file=sys.stderr)
        return EXIT_INTERNAL


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="html2pptx-cli",
        description="受限 HTML → 原生可编辑 pptx",
    )
    sub = parser.add_subparsers(dest="command", required=True)

    validate = sub.add_parser("validate", help="只校验，不生成")
    validate.add_argument("files", nargs="+", help="HTML 文件路径")
    validate.add_argument("--json", action="store_true", help="输出机读报告")
    validate.add_argument("--strict", action="store_true", help="警告也按失败处理")
    validate.add_argument("--lossy-fixes", choices=("apply", "warn"), default="apply")

    build = sub.add_parser("build", help="生成 pptx")
    build.add_argument("files", nargs="+", help="HTML 文件路径，多份按顺序拼页")
    build.add_argument("-o", "--output", required=True, help="输出 .pptx 路径")
    build.add_argument("--report", help="把 JSON 报告写到该路径")
    build.add_argument("--strict", action="store_true", help="警告也按失败处理")
    build.add_argument("--lossy-fixes", choices=("apply", "warn"), default="apply")
    build.add_argument("--font-strategy", choices=("auto", "portable", "windows"), default="auto")
    build.add_argument("--base-dir", help="图片等资源的根目录（默认取首个 HTML 所在目录）")
    build.add_argument("--no-embed-fonts", action="store_true", help="不内嵌字体子集（默认内嵌）")

    contract = sub.add_parser("contract", help="导出契约")
    contract.add_argument("--format", choices=("json", "md", "llm"), default="md")

    inspect = sub.add_parser("inspect", help="回读 pptx 的形状与几何")
    inspect.add_argument("file", help=".pptx 路径")
    inspect.add_argument("--json", action="store_true", help="输出机读摘要")
    return parser


def _options(args: argparse.Namespace, base_dir: Path) -> CompileOptions:
    return CompileOptions(
        base_dir=base_dir,
        lossy_fixes=args.lossy_fixes,
        fail_on_warnings=args.strict,
        font_strategy=getattr(args, "font_strategy", "portable"),
        embed_fonts=not getattr(args, "no_embed_fonts", False),
    )


def _validate(args: argparse.Namespace) -> int:
    paths = [Path(item) for item in args.files]
    sources = [path.read_text(encoding="utf-8") for path in paths]
    report = validate_html(sources, options=_options(args, paths[0].parent))
    if args.json:
        print(json.dumps(report.to_dict(), ensure_ascii=False, indent=2))
    else:
        _print_human(report, [path.name for path in paths])
    return EXIT_OK if _passed(report, args.strict) else EXIT_CONTRACT


def _build(args: argparse.Namespace) -> int:
    paths = [Path(item) for item in args.files]
    sources = [path.read_text(encoding="utf-8") for path in paths]
    base_dir = Path(args.base_dir) if args.base_dir else paths[0].parent
    result = compile_html(sources, options=_options(args, base_dir))
    if args.report:
        target = Path(args.report)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(
            json.dumps(result.report.to_dict(), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
            newline="\n",
        )
    if not result.ok:
        _print_human(result.report, [path.name for path in paths])
        return EXIT_CONTRACT
    output = result.save(args.output)
    print(
        f"已生成 {output}（{len(result.pptx or b'')} 字节，{len(result.document.slides)} 页，"
        f"{result.report.stats.get('elements', 0)} 个元素）"
    )
    if result.report.fixes:
        print(f"自动修复 {len(result.report.fixes)} 条；警告 {len(result.report.warnings)} 条")
    return EXIT_OK


def _contract(args: argparse.Namespace) -> int:
    contract = load_contract()
    if args.format == "json":
        print(render_json(contract), end="")
    elif args.format == "llm":
        print(render_llm(contract), end="")
    else:
        print(render_markdown(contract), end="")
    return EXIT_OK


def _inspect(args: argparse.Namespace) -> int:
    from pptx import Presentation

    presentation = Presentation(args.file)
    slides: list[dict[str, Any]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        shapes = []
        for shape in slide.shapes:
            text = shape.text_frame.text if shape.has_text_frame else ""
            shapes.append(
                {
                    "type": str(shape.shape_type),
                    "name": shape.name,
                    "box_emu": [shape.left, shape.top, shape.width, shape.height],
                    "text": text[:40],
                }
            )
        slides.append({"slide": index, "shapes": shapes})
    if args.json:
        print(json.dumps({"file": args.file, "slides": slides}, ensure_ascii=False, indent=2))
        return EXIT_OK
    print(f"{args.file}：{len(slides)} 页")
    for slide in slides:
        print(f"  第 {slide['slide']} 页：{len(slide['shapes'])} 个形状")
        for shape in slide["shapes"]:
            preview = f" “{shape['text']}”" if shape["text"] else ""
            print(f"    {shape['type']:<24} {shape['box_emu']}{preview}")
    return EXIT_OK


def _passed(report: Report, strict: bool) -> bool:
    return report.ok and not (strict and report.warnings)


def _print_human(report: Report, names: Sequence[str]) -> None:
    print(f"校验 {', '.join(names)}")
    if report.fixes:
        print(f"自动修复 {len(report.fixes)} 条：")
        for fix in report.fixes:
            lossy = "" if fix.lossless else "（有损）"
            print(
                f"  L{fix.line} [{fix.rule}]{lossy} {fix.prop}: {fix.value_from} → {fix.value_to}"
            )
    if report.warnings:
        print(f"警告 {len(report.warnings)} 条：")
        for warning in report.warnings:
            print(f"  L{warning.line} [{warning.code}] {warning.message}")
            print(f"       hint: {warning.hint}")
    if report.errors:
        print(f"错误 {len(report.errors)} 条：")
        for error in report.errors:
            print(f"  L{error.line} [{error.code}] {error.message}")
            print(f"       hint: {error.hint}")
    else:
        print("错误 0 条")


if __name__ == "__main__":  # pragma: no cover - 便于 python -m 调试
    sys.exit(main())

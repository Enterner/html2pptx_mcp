"""pptx → 受限 HTML 逆转：把参考页版式还原成契约子集，供复用与再生成。

方向与编译器相反：DrawingML 是契约的超集，逆转只取可等价映射的部分，
其余降级并逐条记录（``ReverseResult.notes``），不静默丢弃。度量换算：
1px = 9525 EMU（契约画布口径），字号 1pt = 4/3 px。

保真验收口径：逆转 HTML 重新编译出的 pptx 与原 pptx 各自渲染后做像素
diff，diff 不再下降即视为还原收敛。两类固有降级不是逆转缺陷：
字体白名单外的字体（如主题字体「等线」）落到策略默认字体；图片默认
「拉伸到框」在契约里没有对应值，用「位图预拉伸 + object-fit:cover」
等价复现。
"""

from __future__ import annotations

import argparse
import colorsys
import contextlib
import html as html_mod
import io
import json
import math
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

EMU_PER_PX = 9525
PT_TO_PX = 4 / 3
# bodyPr 内边距的 OOXML 默认值：PPT 未写 lIns/tIns/rIns/bIns 时按此渲染
_DEFAULT_INSETS = (91440, 45720, 91440, 45720)
# 主题 lnStyleLst 三档线宽（EMU）——lnRef idx 1..3 的标准 Office 主题口径
_LN_STYLE_WIDTHS = (6350, 12700, 19050)


def _px(emu: int | None) -> float:
    return round((emu or 0) / EMU_PER_PX, 2)


def _css(v: float) -> str:
    text = f"{v:.2f}".rstrip("0").rstrip(".")
    return text if text else "0"


@dataclass
class Note:
    """一条降级/近似记录：逆转不是无损的，规则与后果必须可追溯。"""

    shape: int
    rule: str
    detail: str


@dataclass
class ReverseResult:
    html_paths: list[Path] = field(default_factory=list)
    asset_paths: list[Path] = field(default_factory=list)
    notes: list[Note] = field(default_factory=list)
    # 页号 → .stage 内片段（不含视口壳）；供调用方按需内联，控制上下文占用
    html_by_page: dict[int, str] = field(default_factory=dict)

    def note(self, shape: int, rule: str, detail: str) -> None:
        self.notes.append(Note(shape, rule, detail))


# ---------------------------------------------------------------- 主题解析
class ThemeResolver:
    """主题解析：schemeClr 名 → #RRGGBB（按母版 clrMap 处理 bg1/tx1 别名），
    以及 minor/major 字体槽——run 未显式指定字体时按 PPT 语义落主题字体。

    颜色变换子节点（lumMod/shade 等）就地应用——模板深浅变体几乎全靠它们。
    """

    def __init__(self, master: Any, result: ReverseResult) -> None:
        self._result = result
        self._scheme: dict[str, str] = {}
        self._map: dict[str, str] = {}
        self.minor_font: str | None = None
        self.major_font: str | None = None
        theme_part = None
        for rel in master.part.rels.values():
            if "theme" in rel.reltype:
                theme_part = rel.target_part
                break
        if theme_part is not None:
            root = etree.fromstring(theme_part.blob)
            scheme = root.find(qn("a:themeElements") + "/" + qn("a:clrScheme"))
            if scheme is not None:
                for child in scheme:
                    name = etree.QName(child).localname
                    srgb = child.find(qn("a:srgbClr"))
                    sysc = child.find(qn("a:sysClr"))
                    if srgb is not None:
                        self._scheme[name] = srgb.get("val")
                    elif sysc is not None:
                        self._scheme[name] = sysc.get("lastClr") or "000000"
            fonts = root.find(qn("a:themeElements") + "/" + qn("a:fontScheme"))
            if fonts is not None:
                # ea 槽为空是国内模板常态：中文字形随 latin 槽（等线）走
                for slot, attr in (("a:minorFont", "minor_font"), ("a:majorFont", "major_font")):
                    el = fonts.find(qn(slot))
                    if el is None:
                        continue
                    ea = el.find(qn("a:ea")).get("typeface") or None
                    latin = el.find(qn("a:latin")).get("typeface") or None
                    setattr(self, attr, ea or latin)
        cm = master.element.find(qn("p:clrMap"))
        if cm is not None:
            self._map = dict(cm.attrib)

    def resolve(self, name: str) -> str:
        # clrMap 把 bg1/tx1 等槽位重定向到真实主题槽（bg1→lt1 是常见配置）
        mapped = self._map.get(name, name)
        return self._scheme.get(mapped, self._scheme.get(name, "000000"))

    def color_of_element(self, el: Any, shape: int) -> str | None:
        """srgbClr/sysClr/schemeClr 元素 → #RRGGBB；颜色变换子节点就地应用。

        变换口径与 PowerPoint 一致：shade/tint 在 RGB 空间乘/插值，
        lumMod/lumOff 在 HSL 亮度上乘/加——模板的深浅变体几乎全靠这四个。
        """
        if el is None:
            return None
        tag = etree.QName(el).localname
        value: str | None
        if tag == "srgbClr":
            value = el.get("val")
        elif tag == "sysClr":
            value = el.get("lastClr") or "000000"
        elif tag == "schemeClr":
            value = self.resolve(el.get("val"))
        else:
            return None
        if not value:
            return None
        r, g, b = (
            int(value[0:2], 16) / 255.0,
            int(value[2:4], 16) / 255.0,
            int(value[4:6], 16) / 255.0,
        )
        for child in el:
            kind = etree.QName(child).localname
            raw = int(child.get("val") or 0) / 100000.0
            if kind == "shade":
                r, g, b = r * raw, g * raw, b * raw
            elif kind == "tint":
                r, g, b = r + (1 - r) * raw, g + (1 - g) * raw, b + (1 - b) * raw
            elif kind in ("lumMod", "lumOff"):
                hue, light, sat = colorsys.rgb_to_hls(r, g, b)
                if kind == "lumMod":
                    light *= raw
                else:
                    light += raw
                r, g, b = colorsys.hls_to_rgb(hue, min(1.0, max(0.0, light)), sat)
            elif kind == "alpha":
                self._result.note(shape, "theme.alpha-dropped", f"alpha={raw:g} 丢弃，按不透明处理")
        clamp = lambda c: round(min(1.0, max(0.0, c)) * 255)  # noqa: E731
        return f"{clamp(r):02X}{clamp(g):02X}{clamp(b):02X}"

    def color_in(self, parent: Any, shape: int) -> str | None:
        """在 fill/ln 等父元素下找第一个颜色子节点。"""
        if parent is None:
            return None
        for child in parent:
            hexv = self.color_of_element(child, shape)
            if hexv is not None:
                return hexv
        return None


# ---------------------------------------------------------------- 几何映射
def _xfrm_px(el: Any) -> tuple[float, float, float, float, float, bool, bool] | None:
    """a:xfrm → (left, top, width, height, rot_deg, flipH, flipV)；无则 None。"""
    if el is None:
        return None
    off = el.find(qn("a:off"))
    ext = el.find(qn("a:ext"))
    if off is None or ext is None:
        return None
    rot = int(el.get("rot") or 0) / 60000.0
    return (
        _px(int(off.get("x"))),
        _px(int(off.get("y"))),
        _px(int(ext.get("cx"))),
        _px(int(ext.get("cy"))),
        rot,
        el.get("flipH") == "1",
        el.get("flipV") == "1",
    )


def _prst_polygon(
    prst: str, adjs: dict[str, int], w: float, h: float
) -> list[tuple[float, float]] | None:
    """预设几何 → clip-path 多边形顶点（0–1 分数）；非多边形映射返回 None。"""
    ss = min(w, h) or 1.0
    if prst == "rtTriangle":
        return [(0.0, 0.0), (0.0, 1.0), (1.0, 1.0)]
    if prst == "triangle":
        return [(0.5, 0.0), (0.0, 1.0), (1.0, 1.0)]
    if prst == "rightArrow":
        a1 = adjs.get("adj1", 50000)
        a2 = adjs.get("adj2", 50000)
        dx2 = ss * a2 / 100000.0
        dy1 = h * a1 / 200000.0
        x2 = max(0.0, (w - dx2) / w)
        return [
            (0.0, (h / 2 - dy1) / h),
            (x2, (h / 2 - dy1) / h),
            (x2, 0.0),
            (1.0, 0.5),
            (x2, 1.0),
            (x2, (h / 2 + dy1) / h),
            (0.0, (h / 2 + dy1) / h),
        ]
    return None


def _round_radius(prst: str, adjs: dict[str, int], w: float, h: float) -> str | None:
    if prst == "ellipse":
        return "50%"
    if prst == "roundRect":
        val = adjs.get("adj", 16667)
        return _css(min(0.5, val / 100000.0) * min(w, h)) + "px"
    return None


def _dash_style(prst_dash: str | None) -> str | None:
    if prst_dash in ("dash", "lgDash", "lgDashDot"):
        return "dashed"
    if prst_dash in ("sysDot", "dot", "lgDashDotDot", "sysDashDot", "dashDot"):
        return "dotted"
    return None


def _adjs_of(sp_pr: Any) -> tuple[str | None, dict[str, int]]:
    el = sp_pr.find(qn("a:prstGeom"))
    if el is None:
        return None, {}
    adjs: dict[str, int] = {}
    for gd in el.findall(qn("a:avLst") + "/" + qn("a:gd")):
        fmla = gd.get("fmla") or ""
        if fmla.startswith("val "):
            with contextlib.suppress(ValueError):
                adjs[gd.get("name")] = int(fmla[4:])
    return el.get("prst"), adjs


# ---------------------------------------------------------------- 文本映射
def _runs_html(
    paragraph: Any,
    palette: ThemeResolver,
    shape: int,
    result: ReverseResult,
    placeholder: bool = False,
) -> str:
    """段落 → 内联标签串；b/i 走标签，颜色/字号/字体走 span（契约内联白名单）。"""
    parts: list[str] = []
    for child in paragraph._p:
        tag = etree.QName(child).localname
        if tag == "br":
            parts.append("<br>")
            continue
        if tag == "fld":
            # 字段（页码等）按纯文本落进内容，样式取默认
            parts.append(html_mod.escape("".join(child.itertext())))
            continue
        if tag != "r":
            continue
        text = "".join(t.text or "" for t in child.findall(qn("a:t")))
        if not text:
            continue
        if placeholder:
            # 只复用格式时的占位：等长全角字，宽度与原文最接近，版式不塌
            text = "文" * len(text)
        r_pr = child.find(qn("a:rPr"))
        bold = italic = False
        color = size_px = font = None
        if r_pr is not None:
            bold = r_pr.get("b") == "1"
            italic = r_pr.get("i") == "1"
            color = palette.color_in(r_pr.find(qn("a:solidFill")), shape)
            latin = r_pr.find(qn("a:latin"))
            if latin is not None:
                font = latin.get("typeface") or None
            # CJK 字形由 ea 槽决定，且模板常只写 ea（latin 留空走主题）——
            # ea 优先于 latin，否则缺字体的回退行为与原版不一致
            ea = r_pr.find(qn("a:ea"))
            if ea is not None:
                font = ea.get("typeface") or font
            sz = r_pr.get("sz")
            if sz:
                size_px = _css(round(int(sz) / 100 * PT_TO_PX, 2))
            if r_pr.get("u") not in (None, "none"):
                result.note(shape, "text.underline-dropped", f"下划线不支持：{text[:12]!r}")
        # PPT 语义：run 无显式字体（或引用主题槽）时落主题字体，不是策略默认
        if font in (None, "", "+mn-lt", "+mn-ea"):
            font = palette.minor_font
        elif font == "+mj-lt" or font == "+mj-ea":
            font = palette.major_font
        span_style = []
        if color:
            span_style.append(f"color:#{color}")
        if size_px:
            span_style.append(f"font-size:{size_px}px")
        if font:
            span_style.append(f"font-family:'{font}'")
        inner = html_mod.escape(text)
        if span_style:
            inner = f'<span style="{";".join(span_style)}">{inner}</span>'
        if italic:
            inner = f"<i>{inner}</i>"
        if bold:
            inner = f"<b>{inner}</b>"
        parts.append(inner)
    return "".join(parts)


def _text_html(
    shape: Any,
    box_w: float,
    box_h: float,
    palette: ThemeResolver,
    shape_idx: int,
    result: ReverseResult,
    placeholder: bool = False,
) -> tuple[str, dict[str, str]]:
    """文本框 → (内嵌 HTML, 样式字典)；含 anchor/对齐/行距/内边距口径。"""
    del box_w  # 换行数估算未用到宽度，保留参数以对齐调用方语义
    tf = shape.text_frame
    body = tf._txBody.find(qn("a:bodyPr"))
    style: dict[str, str] = {}

    ins = (
        int(body.get("lIns") or _DEFAULT_INSETS[0]),
        int(body.get("tIns") or _DEFAULT_INSETS[1]),
        int(body.get("rIns") or _DEFAULT_INSETS[2]),
        int(body.get("bIns") or _DEFAULT_INSETS[3]),
    )
    # 编译器默认内边距是 0，与 PPT 默认差 9.6/4.8px，必须显式写全避免漂移
    style["padding"] = (
        f"{_css(_px(ins[1]))}px {_css(_px(ins[2]))}px {_css(_px(ins[3]))}px {_css(_px(ins[0]))}px"
    )

    if body.get("wrap") == "none":
        style["white-space"] = "nowrap"
    if body.get("vert") == "eaVert":
        style["writing-mode"] = "vertical-rl"
    elif body.get("vert") not in (None, "horz"):
        result.note(shape_idx, "text.vert-unsupported", f"vert={body.get('vert')} 按横排处理")

    fit = body.find(qn("a:normAutofit"))
    if fit is not None and fit.get("fontScale"):
        result.note(
            shape_idx,
            "text.autofit-baked",
            f"normAutofit fontScale={fit.get('fontScale')} 已按缩放后字号处理",
        )
    if body.find(qn("a:spAutoFit")) is not None:
        result.note(shape_idx, "text.autofit-shape", "spAutoFit（框随文变）按声明盒处理")

    paras = tf.paragraphs
    if not paras:
        return "", style

    first_ppr = paras[0]._p.find(qn("a:pPr"))
    algn = first_ppr.get("algn") if first_ppr is not None else None
    if algn == "ctr":
        style["text-align"] = "center"
    elif algn == "r":
        style["text-align"] = "right"
    elif algn == "just":
        style["text-align"] = "justify"
    for idx, para in enumerate(paras):
        ppr = para._p.find(qn("a:pPr"))
        this_algn = ppr.get("algn") if ppr is not None else None
        if this_algn != algn:
            result.note(
                shape_idx,
                "text.align-per-paragraph",
                f"第 {idx + 1} 段对齐 {this_algn} 与首段 {algn} 不同，按首段",
            )

    anchor = body.get("anchor")
    # 段级字距：契约把 letter-spacing 放在段落容器上，取首处 spc 整段统一；
    # span 内联白名单不含字距，run 级差异化字距无法表达。
    for para in paras:
        for r_el in para._p.findall(qn("a:r")):
            r_pr = r_el.find(qn("a:rPr"))
            if r_pr is not None and r_pr.get("spc"):
                spc_px = int(r_pr.get("spc")) / 100 * PT_TO_PX
                style["letter-spacing"] = _css(spc_px) + "px"
                break
        else:
            continue
        break

    # 行距：显式 lnSpc 优先；anchor=ctr 的单段文本走「行高=盒高」居中惯用法；
    # 多段锚中按均分行高近似，行数由编译器包装结果决定，可能差半行。
    ppr0 = paras[0]._p.find(qn("a:pPr"))
    ln_spc = ppr0.find(qn("a:lnSpc")) if ppr0 is not None else None
    explicit_lh: str | None = None
    if ln_spc is not None:
        pct = ln_spc.find(qn("a:spcPct"))
        pts = ln_spc.find(qn("a:spcPts"))
        if pct is not None:
            multiple = int(pct.get("val")) / 100000.0
            if 0 <= multiple <= 5:
                explicit_lh = _css(multiple)
            else:
                explicit_lh = None
                result.note(
                    shape_idx, "text.line-height-clamped", f"lnSpc {multiple:g} 超契约倍数区间"
                )
        elif pts is not None:
            explicit_lh = _css(_px(int(pts.get("val")))) + "px"
    if explicit_lh:
        style["line-height"] = explicit_lh
    elif anchor == "ctr":
        if len(paras) == 1:
            style["line-height"] = _css(box_h) + "px"
        else:
            style["line-height"] = _css(round(box_h / len(paras), 2)) + "px"
            result.note(
                shape_idx, "text.anchor-center-approx", f"{len(paras)} 段按均分行高近似垂直居中"
            )
    elif anchor == "b":
        result.note(shape_idx, "text.anchor-bottom-unsupported", "anchor=b 无契约对应，按顶端处理")

    html_parts = [_runs_html(para, palette, shape_idx, result, placeholder) for para in paras]
    return "<br>".join(html_parts), style


# ---------------------------------------------------------------- 主遍历
class _SlideReverser:
    """单页逆转：按文档序展平形状树（含组合），产出内联样式元素序列。"""

    def __init__(
        self,
        slide: Any,
        palette: ThemeResolver,
        assets_dir: Path,
        result: ReverseResult,
        page_label: str,
        placeholder: bool = False,
        top_selected: set[int] | None = None,
    ) -> None:
        self.slide = slide
        self.palette = palette
        self.assets_dir = assets_dir
        self.result = result
        self.page_label = page_label
        self.placeholder = placeholder
        # 顶层形状索引白名单；None = 全部。组合被选中时整棵子树跟随。
        self.top_selected = top_selected
        self.body: list[str] = []
        self.shape_seq = 0
        self.top_selected = top_selected

    # ---- 填充/描边/效果 ----
    def _fill_html(self, sp_pr: Any, style_el: Any, shape_idx: int) -> dict[str, str]:
        style: dict[str, str] = {}
        if sp_pr is not None:
            if sp_pr.find(qn("a:noFill")) is not None:
                return style
            solid = sp_pr.find(qn("a:solidFill"))
            if solid is not None:
                hexv = self.palette.color_in(solid, shape_idx)
                if hexv:
                    style["background-color"] = f"#{hexv}"
                    return style
            grad = sp_pr.find(qn("a:gradFill"))
            if grad is not None:
                style.update(self._gradient_html(grad, shape_idx))
                return style
        if style_el is not None:
            fill_ref = style_el.find(qn("a:fillRef"))
            if fill_ref is not None and int(fill_ref.get("idx") or 0) > 0:
                hexv = self.palette.color_of_element(fill_ref.find(qn("a:schemeClr")), shape_idx)
                if hexv:
                    style["background-color"] = f"#{hexv}"
        return style

    def _gradient_html(self, grad: Any, shape_idx: int) -> dict[str, str]:
        stops: list[tuple[float, str]] = []
        for gs in grad.findall(qn("a:gsLst") + "/" + qn("a:gs")):
            pos = int(gs.get("pos") or 0) / 1000.0
            hexv = self.palette.color_in(gs, shape_idx)
            if hexv is not None:
                stops.append((pos, hexv))
        stops.sort()
        lin = grad.find(qn("a:lin"))
        if lin is not None:
            # CSS 角度 0=向上顺时针；OOXML ang 0=向右顺时针 → css = ooxml + 90
            ang = int(lin.get("ang") or 0) / 60000.0
            css_ang = round((ang + 90.0) % 360.0)
        else:
            css_ang = 180
            self.result.note(shape_idx, "fill.gradient-path", "path 渐变按 180deg 线性近似")
        if len(stops) > 2:
            self.result.note(shape_idx, "fill.gradient-2-stops", f"{len(stops)} 色标截为首尾两枚")
            stops = [stops[0], stops[-1]]
        if len(stops) < 2:
            return {}
        a, b = stops
        value = f"linear-gradient({css_ang}deg, #{a[1]} {_css(a[0])}%, #{b[1]} {_css(b[0])}%)"
        return {"background": value}

    def _line_style(self, sp_pr: Any, style_el: Any, shape_idx: int) -> dict[str, str]:
        style: dict[str, str] = {}
        ln = sp_pr.find(qn("a:ln")) if sp_pr is not None else None
        if ln is not None:
            if ln.find(qn("a:noFill")) is not None:
                return style
            hexv = self.palette.color_in(ln.find(qn("a:solidFill")), shape_idx)
            w = ln.get("w")
            width_px = _px(int(w)) if w else 1.0
            if hexv:
                dash_el = ln.find(qn("a:prstDash"))
                dash = _dash_style(dash_el.get("val") if dash_el is not None else None)
                dash_part = f" {dash}" if dash else ""
                style["border"] = f"{_css(width_px)}px{dash_part} solid #{hexv}"
                return style
        if style_el is not None:
            ln_ref = style_el.find(qn("a:lnRef"))
            if ln_ref is not None and int(ln_ref.get("idx") or 0) > 0:
                idx = min(int(ln_ref.get("idx")), 3)
                width_px = _px(_LN_STYLE_WIDTHS[idx - 1])
                hexv = self.palette.color_of_element(ln_ref.find(qn("a:schemeClr")), shape_idx)
                if hexv:
                    style["border"] = f"{_css(width_px)}px solid #{hexv}"
                    self.result.note(shape_idx, "line.style-ref", "lnRef 主题线型按 idx 档位近似")
        return style

    def _shadow_html(self, sp_pr: Any, style_el: Any, shape_idx: int) -> dict[str, str]:
        style: dict[str, str] = {}
        eff = sp_pr.find(qn("a:effectLst")) if sp_pr is not None else None
        if eff is not None:
            shdw = eff.find(qn("a:outerShdw"))
            if shdw is not None:
                blur = _px(int(shdw.get("blurRad") or 0))
                dist = _px(int(shdw.get("dist") or 0))
                rad = math.radians(int(shdw.get("dir") or 0) / 60000.0)
                dx = round(dist * math.cos(rad), 2)
                dy = round(dist * math.sin(rad), 2)
                hexv = self.palette.color_in(shdw, shape_idx) or "000000"
                alpha_el = shdw.find(".//" + qn("a:alpha"))
                alpha = int(alpha_el.get("val")) / 100000.0 if alpha_el is not None else 1.0
                rgb = f"{int(hexv[0:2], 16)},{int(hexv[2:4], 16)},{int(hexv[4:6], 16)}"
                style["box-shadow"] = (
                    f"{_css(dx)}px {_css(dy)}px {_css(blur)}px rgba({rgb},{_css(alpha)})"
                )
                return style
        if style_el is not None:
            effect_ref = style_el.find(qn("a:effectRef"))
            if effect_ref is not None and int(effect_ref.get("idx") or 0) > 0:
                self.result.note(
                    shape_idx, "effect.style-ref", "effectRef 主题效果无契约对应，丢弃"
                )
        return style

    # ---- 图片 ----
    def _picture(self, sh: Any, x: float, y: float, w: float, h: float, shape_idx: int) -> str:
        image = sh.image
        blob = image.blob
        src_rect = sh._element.find(qn("p:blipFill") + "/" + qn("a:srcRect"))
        if src_rect is not None:
            src = Image.open(io.BytesIO(blob))
            sw, sh_h = src.size
            crop = (
                round(sw * int(src_rect.get("l") or 0) / 100000.0),
                round(sh_h * int(src_rect.get("t") or 0) / 100000.0),
                round(sw * (1 - int(src_rect.get("r") or 0) / 100000.0)),
                round(sh_h * (1 - int(src_rect.get("b") or 0) / 100000.0)),
            )
            blob = src.crop(crop)
        # PPT 图片默认「拉伸到框」；契约无 stretch，用位图预拉伸 + cover 等价复现：
        # 位图先按盒比例重采样（2x 保清晰度），cover 落框后与拉伸逐像素等价
        src_img = Image.open(io.BytesIO(blob)).convert("RGB")
        src_img = src_img.resize((max(1, round(w * 2)), max(1, round(h * 2))))
        buffer = io.BytesIO()
        src_img.save(buffer, format="PNG")

        self.shape_seq += 1
        name = f"{self.page_label}_{self.shape_seq:02d}.png"
        path = self.assets_dir / name
        path.write_bytes(buffer.getvalue())
        self.result.asset_paths.append(path)
        src_attr = f"{self.assets_dir.name}/{name}"
        return (
            f'<img src="{src_attr}" '
            f'style="position:absolute;left:{_css(x)}px;top:{_css(y)}px;'
            f'width:{_css(w)}px;height:{_css(h)}px;object-fit:cover;box-sizing:border-box">'
        )

    # ---- 连接线 ----
    def _connector(
        self,
        sh: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        flip_h: bool,
        flip_v: bool,
        shape_idx: int,
    ) -> str:
        sp_pr = sh._element.find(qn("p:spPr"))
        ln = sp_pr.find(qn("a:ln"))
        hexv = (
            self.palette.color_in(ln.find(qn("a:solidFill")), shape_idx) if ln is not None else None
        )
        w_emu = ln.get("w") if ln is not None else None
        width_px = _px(int(w_emu)) if w_emu else 1.0
        dash_el = ln.find(qn("a:prstDash")) if ln is not None else None
        dash = _dash_style(dash_el.get("val") if dash_el is not None else None)
        x1, x2 = (w, 0.0) if flip_h else (0.0, w)
        y1, y2 = (h, 0.0) if flip_v else (0.0, h)
        if dash == "dashed":
            dash_attr = ' stroke-dasharray="8 4"'
        elif dash == "dotted":
            dash_attr = ' stroke-dasharray="2 3"'
        else:
            dash_attr = ""
        stroke = f"#{hexv or '000000'}"
        return (
            f'<svg style="position:absolute;left:{_css(x)}px;top:{_css(y)}px" '
            f'width="{_css(w)}" height="{_css(h)}">'
            f'<line x1="{_css(x1)}" y1="{_css(y1)}" x2="{_css(x2)}" y2="{_css(y2)}" '
            f'stroke="{stroke}" stroke-width="{_css(width_px)}"{dash_attr}></line></svg>'
        )

    # ---- 形状 ----
    def _shape_div(
        self,
        sh: Any,
        el: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        shape_idx: int,
        flip_h: bool = False,
        flip_v: bool = False,
    ) -> str:
        sp_pr = el.find(qn("p:spPr"))
        style_el = el.find(qn("p:style"))
        prst, adjs = _adjs_of(sp_pr)
        prst = prst or ""

        style = self._fill_html(sp_pr, style_el, shape_idx)
        style.update(self._line_style(sp_pr, style_el, shape_idx))
        style.update(self._shadow_html(sp_pr, style_el, shape_idx))

        radius = _round_radius(prst, adjs, w, h)
        if radius:
            style["border-radius"] = radius
        polygon = _prst_polygon(prst, adjs, w, h) if radius is None else None
        if polygon:
            if flip_h or flip_v:
                # flip 在顶点上镜像（rect/roundRect/ellipse 对称无视觉差，无需处理）；
                # rot 仍由 transform rotate 承担，两者叠加与 PPT 的几何语义一致
                polygon = [
                    (1.0 - px if flip_h else px, 1.0 - py if flip_v else py) for px, py in polygon
                ]
            pts = ", ".join(f"{_css(px * 100)}% {_css(py * 100)}%" for px, py in polygon)
            style["clip-path"] = f"polygon({pts})"
        known = ("rect", "roundRect", "ellipse", "rtTriangle", "triangle", "rightArrow")
        if prst is not None and prst not in known:
            self.result.note(shape_idx, "geometry.preset", f"prst={prst} 按矩形处理")

        xfrm = el.find(".//" + qn("a:xfrm"))
        deg = int(xfrm.get("rot") or 0) / 60000.0
        if deg:
            style["transform"] = f"rotate({_css(deg)}deg)"

        inner = ""
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            inner, text_style = _text_html(
                sh, w, h, self.palette, shape_idx, self.result, self.placeholder
            )
            style.update(text_style)

        parts = [
            f"position:absolute;left:{_css(x)}px;top:{_css(y)}px;"
            f"width:{_css(w)}px;height:{_css(h)}px",
            # PPT 形状宽高是「边到边」；CSS 默认 content-box 会把 padding/border
            # 加在外面（pill 实测多 19px）。契约有无损 absorb 规则，编译时吸收。
            "box-sizing:border-box",
        ]
        parts += [f"{key}:{value}" for key, value in style.items()]
        return f'<div style="{";".join(parts)}">{inner}</div>'

    # ---- 树遍历 ----
    def _walk(
        self,
        shapes: Any,
        dx: float = 0.0,
        dy: float = 0.0,
        sx: float = 1.0,
        sy: float = 1.0,
        base_idx: int = 0,
    ) -> None:
        for offset, sh in enumerate(shapes):
            shape_idx = base_idx + offset
            # 元素选择只在顶层过滤；组合选中则整棵子树跟随
            if base_idx == 0 and self.top_selected is not None and offset not in self.top_selected:
                continue
            el = sh._element
            xfrm = _xfrm_px(el.find(".//" + qn("a:xfrm")))
            if xfrm is None:
                self.result.note(shape_idx, "geometry.no-xfrm", "无 xfrm，跳过")
                continue
            x, y, w, h, rot, flip_h, flip_v = xfrm
            x, y = x * sx + dx, y * sy + dy
            w, h = w * sx, h * sy

            if etree.QName(el).localname == "grpSp":
                gx = el.find(qn("p:grpSpPr") + "/" + qn("a:xfrm"))
                ch_off = gx.find(qn("a:chOff"))
                ch_ext = gx.find(qn("a:chExt"))
                cox, coy = _px(int(ch_off.get("x"))), _px(int(ch_off.get("y")))
                cex = _px(int(ch_ext.get("cx"))) or 1.0
                cey = _px(int(ch_ext.get("cy"))) or 1.0
                nsx = sx * (w / cex) if cex else sx
                nsy = sy * (h / cey) if cey else sy
                ndx = x - cox * nsx
                ndy = y - coy * nsy
                if rot or gx.get("flipH") == "1" or gx.get("flipV") == "1":
                    self.result.note(
                        shape_idx, "group.transform", "组合带旋转/翻转，按展平轴对齐盒处理"
                    )
                self._walk(sh.shapes, ndx, ndy, nsx, nsy, shape_idx * 100)
                continue

            if etree.QName(el).localname == "pic":
                self.body.append(self._picture(sh, x, y, w, h, shape_idx))
                continue
            if etree.QName(el).localname == "cxnSp":
                self.body.append(self._connector(sh, x, y, w, h, flip_h, flip_v, shape_idx))
                continue
            self.body.append(self._shape_div(sh, el, x, y, w, h, shape_idx, flip_h, flip_v))

    # ---- 页面 ----
    def render(self) -> str:
        c_sld = self.slide._element.find(qn("p:cSld"))
        bg = c_sld.find(qn("p:bg"))
        bg_style = "#FFFFFF"
        if bg is not None:
            bg_pr = bg.find(qn("p:bgPr"))
            if bg_pr is not None:
                solid = bg_pr.find(qn("a:solidFill"))
                hexv = self.palette.color_in(solid, 0) if solid is not None else None
                if hexv:
                    bg_style = f"#{hexv}"
        self._walk(self.slide.shapes)
        return (
            '<div style="position:absolute;left:0;top:0;width:1280px;height:720px;'
            f'background:{bg_style}">' + "".join(self.body) + "</div>"
        )


_SHELL = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<title>{title}</title>
<style>
  html, body {{ margin:0; padding:0; width:100%; height:100%; }}
  body {{ background:#000000; font-family:"Microsoft YaHei","微软雅黑",sans-serif;
          overflow:hidden; }}
  .viewport {{ width:100vw; height:100vh; display:flex; align-items:center;
               justify-content:center; }}
  .stage {{ position:relative; width:1280px; height:720px; overflow:hidden;
            flex:0 0 auto; }}
</style>
<!--
{notes}
-->
</head>
<body>
<div class="viewport">
  <div class="stage">
{body}
  </div>
</div>
</body>
</html>
"""


def reverse_pptx(
    pptx_path: str | Path,
    out_dir: str | Path,
    *,
    pages: list[int] | None = None,
    elements: dict[int, list[int]] | None = None,
    placeholder_text: bool = False,
) -> ReverseResult:
    """逆转入口。

    ``pages``：1 起始页号白名单，None = 全部；``elements``：页号 → 顶层形状
    索引白名单（组合选中则整棵子树跟随），None = 全部——两级粒度支撑
    「整册全逆转」到「只抽标题样式」的连续谱。
    """
    pptx_path = Path(pptx_path)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    assets_dir = out_dir / "assets"
    assets_dir.mkdir(exist_ok=True)

    presentation = Presentation(str(pptx_path))
    result = ReverseResult()
    palette = ThemeResolver(presentation.slide_masters[0], result)

    page_set = set(pages) if pages is not None else None
    element_map = {int(k): set(v) for k, v in elements.items()} if elements else {}

    stem = pptx_path.stem
    many = len(presentation.slides) > 1
    for page, slide in enumerate(presentation.slides, 1):
        if page_set is not None and page not in page_set:
            continue
        label = f"{stem}_p{page:02d}" if many else stem
        reverser = _SlideReverser(
            slide,
            palette,
            assets_dir,
            result,
            label,
            placeholder=placeholder_text,
            top_selected=element_map.get(page),
        )
        body = reverser.render()
        notes = "\n".join(f"  [{n.shape}] {n.rule}: {n.detail}" for n in result.notes)
        html = _SHELL.format(title=f"{stem} p{page}", notes=notes or "  (无降级)", body=body)
        name = f"{stem}_p{page:02d}.html" if many else f"{stem}.html"
        path = out_dir / name
        path.write_text(html, encoding="utf-8")
        result.html_paths.append(path)
        result.html_by_page[page] = body
    return result


def _guess_role(
    tag: str,
    box: tuple[int, ...],
    has_text: bool,
    is_max_font: bool,
    canvas: tuple[int, int],
) -> str:
    """角色初判：给 agent 的语义线索，位置/尺寸/字级的启发式，最终由 agent 定夺。"""
    _left, top, w, h = box
    cw, ch = canvas
    area = w * h
    if tag == "picture":
        return "picture"
    if tag == "line":
        return "line"
    if tag == "group":
        return "group"
    if tag == "table":
        return "table"
    if area >= 0.95 * cw * ch:
        return "bg"
    if w >= cw * 0.9 and h <= ch * 0.22:
        return "header" if top <= ch * 0.18 else "banner"
    if h <= 4 or w <= 4:
        return "divider"
    if has_text:
        if is_max_font and top <= ch * 0.45:
            return "title"
        if w <= cw * 0.45 and h <= ch * 0.2:
            return "badge"
        return "text"
    return "decoration"


def outline_pptx(pptx_path: str | Path) -> dict[str, Any]:
    """整册盘点：主题色板/字体体系 + 每页形状清单与角色初判（紧凑 JSON）。

    这是选择性逆转的「选片依据」：agent 扫清单定哪些是格式，再调
    ``reverse_pptx`` 按页/按元素抽取，避免整册 HTML 灌进上下文。
    """
    pptx_path = Path(pptx_path)
    presentation = Presentation(str(pptx_path))
    result = ReverseResult()
    resolver = ThemeResolver(presentation.slide_masters[0], result)
    canvas = (round(_px(presentation.slide_width)), round(_px(presentation.slide_height)))

    pages_out: list[dict[str, Any]] = []
    for page, slide in enumerate(presentation.slides, 1):
        # 页内最大字号 + 位置靠上者判为标题
        texts: dict[int, dict[str, Any]] = {}
        max_font = 0.0
        for i, sh in enumerate(slide.shapes):
            if not getattr(sh, "has_text_frame", False):
                continue
            runs = [run for para in sh.text_frame.paragraphs for run in para.runs]
            if not runs:
                continue
            sizes = [run.font.size.pt for run in runs if run.font.size is not None]
            fonts = {run.font.name for run in runs if run.font.name}
            texts[i] = {
                "pt": round(max(sizes), 1) if sizes else None,
                "fonts": sorted(fonts),
            }
            if sizes:
                max_font = max(max_font, max(sizes))

        shapes_out: list[dict[str, Any]] = []
        roles: dict[str, int] = {}
        for i, sh in enumerate(slide.shapes):
            el = sh._element
            tag = etree.QName(el).localname
            kind = {
                "sp": "shape",
                "grpSp": "group",
                "pic": "picture",
                "cxnSp": "line",
                "graphicFrame": "table",
            }.get(tag, tag)
            xfrm = _xfrm_px(el.find(".//" + qn("a:xfrm")))
            box = [round(xfrm[0]), round(xfrm[1]), round(xfrm[2]), round(xfrm[3])] if xfrm else None
            text_info = texts.get(i)
            has_text = bool(sh.has_text_frame and sh.text_frame.text.strip())
            is_title = False
            if text_info and text_info["pt"] and text_info["pt"] == max_font:
                top = box[1] if box else 0
                if top <= canvas[1] * 0.45:
                    is_title = True
            role = ""
            if box:
                role = _guess_role(kind, tuple(box), has_text, is_title, canvas)
            entry: dict[str, Any] = {"i": i, "t": kind, "role": role}
            if box:
                entry["box"] = box
            if has_text:
                text = sh.text_frame.text.replace("\n", "|")
                entry["text"] = text[:14] + ("…" if len(text) > 14 else "")
                if text_info and text_info["pt"]:
                    entry["pt"] = text_info["pt"]
                if text_info and text_info["fonts"]:
                    entry["font"] = text_info["fonts"][0]
            shapes_out.append(entry)
            roles[role or kind] = roles.get(role or kind, 0) + 1
        pages_out.append(
            {"page": page, "shapes": len(shapes_out), "roles": roles, "items": shapes_out}
        )

    return {
        "file": str(pptx_path),
        "pages": len(presentation.slides),
        "canvas": {"w": canvas[0], "h": canvas[1]},
        "theme": {
            "colors": {
                "bg1": resolver.resolve("bg1"),
                "tx1": resolver.resolve("tx1"),
                "accent1": resolver.resolve("accent1"),
                "accent2": resolver.resolve("accent2"),
                "accent3": resolver.resolve("accent3"),
            },
            "fonts": {"minor": resolver.minor_font, "major": resolver.major_font},
        },
        "inventory": pages_out,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="参考 pptx 逆转为受限 HTML（契约子集）")
    parser.add_argument("pptx", type=Path, help="输入 pptx")
    parser.add_argument("-o", "--out", type=Path, required=True, help="输出目录")
    parser.add_argument("--pages", type=str, default=None, help="页号白名单，如 1,3；缺省全部")
    args = parser.parse_args(argv)

    pages = [int(x) for x in args.pages.split(",")] if args.pages else None
    result = reverse_pptx(args.pptx, args.out, pages=pages)
    for path in result.html_paths:
        print(f"HTML  {path}")
    print(f"资产  {len(result.asset_paths)} 个 -> {args.out / 'assets'}")
    if result.notes:
        print("降级记录：")
        for note in result.notes:
            print(f"  [{note.shape}] {note.rule}: {note.detail}")
    else:
        print("降级记录：无")
    report_path = args.out / "reverse_report.json"
    report_path.write_text(
        json.dumps(
            {
                "html": [str(p) for p in result.html_paths],
                "assets": [str(p) for p in result.asset_paths],
                "notes": [vars(n) for n in result.notes],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"报告  {report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
